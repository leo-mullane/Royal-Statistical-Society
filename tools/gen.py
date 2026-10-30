"""Generate the RSS prototype's src/ pages, data.js, contents hub, module
library and requirements trace. Run from the project root:

    python3 tools/gen.py && node build-includes.js && node validate.js
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *  # noqa
import pages_a, pages_b, pages_c, pages_d  # noqa

VERSION = 'v0.2'

# Contents hub: grouped by the proposed top-level navigation, so the hub doubles as the sitemap
SITEMAP = [
    ('Global', [
        ('home.html', 'Home', 'Task-led home page, personalised for members, with events, training, news carousel and the RSS family', 'REQ-101, REQ-238, REQ-237, REQ-239, REQ-240'),
        ('search.html?q=chartered+statistician', 'Search results', 'Facets with counts, synonyms, typo tolerance, document text, AI summary', 'REQ-037 to REQ-046'),
        ('search-advanced.html', 'Advanced search', 'Opens pre-filled from the search you came from', 'REQ-039'),
        ('basket.html', 'Basket', 'Several purchases in one payment, with promo codes', 'REQ-094, REQ-087'),
        ('sign-in.html', 'Sign in and create an account', 'MFA, registration, password reset, staff sign-in', 'REQ-074 to REQ-080'),
    ]),
    ('Membership', [
        ('membership.html', 'Membership (section landing)', 'Benefits, grades, FAQs', 'REQ-187, REQ-104'),
        ('grades-fees.html', 'Grades and fees', 'Content page with left-hand menu; fees and concessions', 'REQ-029, REQ-189'),
        ('join.html', 'Join (multi-step form)', 'Grade, details, journals, payment by card, Direct Debit or invoice', 'REQ-186, REQ-082, REQ-084, REQ-086'),
    ]),
    ('Professional development', [
        ('professional-development.html', 'Professional development hub', 'Use-based hub: status, CPD, accreditation, courses', 'REQ-034'),
        ('chartered-statistician.html', 'Chartered Statistician (content page)', 'Standard content page template', 'REQ-029, REQ-013'),
        ('accreditation.html', 'Accreditation', 'Expression of interest form', 'REQ-190, REQ-191'),
    ]),
    ('Events and training', [
        ('events-training.html', 'Events and training hub', 'Hub with pre-filtered views, conference, recordings', 'REQ-034, REQ-211, REQ-107'),
        ('events.html', 'Find an event or course', 'One finder for events and training; list and calendar views', 'REQ-210, REQ-211, REQ-212, REQ-031'),
        ('event.html', 'Event page', 'One page per event, tabs, a booking button per format', 'REQ-201, REQ-202, REQ-209'),
        ('course.html', 'Course page', 'Every date and price on one page, waiting list', 'REQ-213, REQ-214, REQ-218'),
        ('book.html', 'Book a course', 'Booking for yourself or others', 'REQ-215, REQ-216, REQ-217'),
        ('conference.html', 'RSS Conference 2027', 'Conference landing page', 'REQ-224 to REQ-233'),
        ('submit-event.html', 'Submit a group event', 'For group officers', 'REQ-203'),
    ]),
    ('Get involved', [
        ('get-involved.html', 'Get involved hub', 'Groups, volunteering, honours, ambassadors', 'REQ-034'),
        ('groups.html', 'Sections and groups', 'Tabs by type, filter by region', 'REQ-102'),
        ('group.html', 'Section or group page', 'Pulls through tagged news and events', 'REQ-102, REQ-163'),
        ('donate.html', 'Donate', 'Gift Aid', 'REQ-090'),
    ]),
    ('Policy and research', [
        ('policy.html', 'Policy and research, with the policy library', 'One filterable library of documents', 'REQ-252'),
        ('campaign.html', 'Campaign or explainer', 'Campaign template', 'REQ-104, REQ-108'),
    ]),
    ('News and publications', [
        ('news.html', 'News', 'Featured story, filters by type and topic', 'REQ-235'),
        ('article.html', 'News article', 'Share buttons, related content', 'REQ-236, REQ-105'),
        ('publications.html', 'Journals and publications', 'OUP access, Significance, Real World Data Science', 'REQ-198, REQ-020, REQ-033'),
    ]),
    ('Resources', [
        ('resources.html', 'Resources hub', 'Audience resource hubs', 'REQ-253'),
        ('for-journalists.html', 'For journalists', 'Audience hub example', 'REQ-253'),
        ('find-an-expert.html', 'Find an expert', 'Professional Register, Consultants and Fellows directories', 'REQ-192, REQ-193, REQ-194'),
        ('jobs.html', 'Jobs', 'Listing, featured posts, advertise', 'REQ-249, REQ-088, REQ-251'),
        ('job.html', 'Job page', 'Description and Apply now', 'REQ-250'),
    ]),
    ('About', [
        ('about.html', 'About the RSS', 'Content page, contact form, security reporting', 'REQ-029, REQ-069, REQ-093'),
    ]),
    ('MyRSS', [
        ('myrss.html', 'Dashboard', 'Member banner, balance and renewal, bookings, publications', 'REQ-173, REQ-179, REQ-181, REQ-188'),
        ('myrss-profile.html', 'Profile and directory listings', 'Sections that save separately; interests', 'REQ-173, REQ-196'),
        ('myrss-bookings.html', 'My events and training', 'Upcoming, past, payment status', 'REQ-180, REQ-208, REQ-219'),
        ('myrss-cpd.html', 'CPD record', 'Add, edit, download', 'REQ-195'),
        ('myrss-groups.html', 'Sections and groups', 'Subscriptions and mailing lists', 'REQ-182'),
        ('myrss-preferences.html', 'Newsletters and preferences', 'Opt in and out', 'REQ-159, REQ-160'),
    ]),
]


def hub():
    groups = ''
    for title, items in SITEMAP:
        lis = ''.join(f'<li><a href="~/pages/{f}">{e(t)}</a> <span class="wf-req">{e(r)}</span><p>{e(s)}</p></li>' for f, t, s, r in items)
        groups += f'<section><h2 class="section-title">{e(title)}</h2><ul class="wf-hub-list">{lis}</ul></section>'
    out = f'''<!DOCTYPE html>
<html lang="en-GB">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>RSS website prototype: contents</title>
  <link rel="stylesheet" href="~/css/base.css">
  <link rel="stylesheet" href="~/css/components.css">
  <link rel="stylesheet" href="~/css/rss.css">
  <link rel="stylesheet" href="~/css/theme.css">
</head>
<body>
  <div class="proto-bar">
    <div class="container proto-bar-inner">
      <a class="proto-bar-home" href="~/index.html">RSS prototype <span>contents</span></a>
      <label class="proto-persona">View as
        <select data-persona-select aria-label="View the site as">
          <option value="guest">Signed out</option>
          <option value="registered">Registered, not a member</option>
          <option value="member">Member (CStat)</option>
          <option value="lapsed">Lapsed member</option>
        </select>
      </label>
      <button type="button" class="notes-toggle" role="switch" aria-checked="false" data-notes-toggle>
        <span class="notes-track" aria-hidden="true"><span class="notes-thumb"></span></span>
        <span class="notes-label">Annotations off</span>
      </button>
    </div>
  </div>
  <main id="main">
    <div class="container">
      <h1>Royal Statistical Society: website prototype</h1>
      <p class="lead">A clickable greyscale prototype of the whole new rss.org.uk, for agreeing structure, navigation and journeys before design. It covers the Must and Should requirements; Could requirements, including ClerksWell's suggestions, are flagged where relevant but not built.</p>
      <p><a class="btn" href="~/pages/home.html">Open the home page</a></p>

      <div class="contents-grid">
        <section>
          <h2 class="section-title">How to review</h2>
          <ul class="wf-hub-list">
            <li><p><strong>Annotations are off.</strong> Switch them on in the dark bar at the top of every page. Yellow notes explain behaviour and CMS capability; red notes are questions for RSS. Each note names the requirement it answers, e.g. <span class="wf-req">REQ-201</span>.</p></li>
            <li><p><strong>View as</strong> (dark bar) switches between signed out, a registered non-member, a member and a lapsed member, so you can see what changes for each.</p></li>
            <li><p>Navigation, search, filters, the calendar, tabs, accordions, forms (with validation), the basket and the cookie banner all work. Individual article and profile links go nowhere by design.</p></li>
            <li><p>Search is fully working. Try "CStat", "bayseian" (a typo) or "census" (found inside a PDF).</p></li>
            <li><p>All people, employers, fees and figures are placeholders.</p></li>
          </ul>
        </section>
        <section>
          <h2 class="section-title">Navigation</h2>
          <ul class="wf-hub-list">
            <li><p><strong>Eight task-based sections</strong> replace the nine organisational ones: Membership, Professional development, Events and training, Get involved, Policy and research, News and publications, Resources, About.</p></li>
            <li><p>The four use-based hubs from RSS's IA document (Professional development, Events and training, Get involved with volunteering, Resources) are top-level sections, not pages under Membership.</p></li>
            <li><p>A utility bar links the RSS family of sites (Significance, Real World Data Science, RSS Publications), MyRSS, the basket and Join.</p></li>
            <li><p>Every mega-menu panel links to its section's landing page, and two panels have search built in (events and experts).</p></li>
          </ul>
          <p><a href="~/modules/library.html">Module library</a>. The requirements trace is in <code>requirements/requirements.md</code> in the repository.</p>
        </section>
      </div>

      <h2 class="section-title">Pages, by section</h2>
      {groups}

      <section>
        <h2 class="section-title">Not in this draft</h2>
        <ul class="wf-hub-list">
          <li><p>Could requirements, including ClerksWell's REQ-276 to REQ-291, and the Won't requirements. Where a Could would change a page, a "Proposed" note says so.</p></li>
          <li><p>Back-office and CMS functions (workflow, permissions, media library, newsletter editor, reporting). These are shown as CMS notes on the pages they affect.</p></li>
          <li><p>The accreditation application portal and the embedded conference programme (Phase 2 in the SoW).</p></li>
        </ul>
      </section>

      <section>
        <h2 class="section-title">Change log</h2>
        <ol class="changelog"><li>v0.1: complete draft from the draft requirements catalogue and the immersion findings.</li><li>v0.2: navigation moves into a menu below 1200px; header shrinks to one 64px row on scroll; event type filter and more discussion meetings; one listing per course; check-your-answers shows everything entered; question on fees and concessions; calendar labels shortened at a word; side menus clear the header.</li></ol>
      </section>
    </div>
  </main>
  <script src="~/js/data.js"></script>
  <script src="~/js/wireframe.js"></script>
  <script src="~/js/rss.js"></script>
</body>
</html>
'''
    with open(os.path.join(SRC, 'index.html'), 'w', encoding='utf-8') as fh:
        fh.write(out)


def library():
    n = Notes()
    body = f'''
      <h1>Module library</h1>
      <p class="lead">Every reusable component, shown once. Pages are composed from these; editors place and configure them in the CMS.</p>
      {n("Components are reusable and editable, placed by drag and drop, each instance with its own content and links.", "REQ-013, REQ-018, REQ-019", "CMS")}
      {n("Page templates are built from these modules, one per content type (news, events, courses, groups, committees, task forces, awards, campaigns, jobs, policy, press releases and more). RSS can clone and edit templates, including header and footer regions.", "REQ-011, REQ-026, REQ-100", "CMS")}
      <h2 class="section-title">Card</h2><div class="module-demo"><div class="grid grid-3"><article class="card"><div class="wf-placeholder ratio-16-9">Image 16:9</div><p class="meta-row"><span class="tag">Type</span><span class="tag">Format</span></p><h3 class="card-title"><a href="#">Card title of a realistic length</a></h3><p class="card-text">Date or summary line</p></article></div></div>
      <h2 class="section-title">Task tile</h2><div class="module-demo"><div class="task-grid"><a class="task" href="#">Find an event or course<span>Short description</span></a></div></div>
      <h2 class="section-title">Listing item</h2><div class="module-demo"><article class="list-item"><div class="date-block"><span class="d">4</span><span class="m">Nov 2026</span></div><div><p class="meta-row"><span class="tag">Discussion meeting</span><span class="tag">Hybrid</span></p><h3><a href="#">Listing item title</a></h3><p class="meta-row"><span>Date and time</span><span>Place</span><span>Price</span></p></div><div class="list-actions"><a class="btn btn-small" href="#">Book</a></div></article></div>
      <h2 class="section-title">Filters and active filter chips</h2><div class="module-demo"><div class="filter-bar"><div class="field"><label for="lib-q">Keyword</label><input type="search" id="lib-q"></div><div class="field"><label for="lib-t">Topic</label><select id="lib-t"><option>All topics</option></select></div></div><div class="active-filters" style="margin-top:1rem"><span class="chip">Online &times;</span><span class="chip">Methodology &times;</span></div></div>
      <h2 class="section-title">Tabs</h2><div class="module-demo"><div class="tabs" role="tablist"><button class="tab active" type="button" data-tab="lib-t1">First</button><button class="tab" type="button" data-tab="lib-t2">Second</button></div><div><div class="tab-panel active" id="lib-t1"><p>First panel</p></div><div class="tab-panel" id="lib-t2"><p>Second panel</p></div></div></div>
      <h2 class="section-title">Accordion</h2><div class="module-demo"><div class="accordion"><div class="accordion-item"><button class="accordion-title" type="button">Question <span class="accordion-icon" aria-hidden="true">&#9662;</span></button><div class="accordion-content"><p>Answer</p></div></div></div></div>
      <h2 class="section-title">Content highlight</h2><div class="module-demo"><div class="content-highlight"><div class="wf-placeholder ratio-16-9">Image 16:9</div><div><h3>Heading</h3><p>Text and a call to action.</p><a class="btn" href="#">Call to action</a></div></div></div>
      <h2 class="section-title">Statistics</h2><div class="module-demo"><div class="statistics"><div class="stat-item"><div class="stat-value">10,000+</div><div class="stat-title">Members</div></div><div class="stat-item"><div class="stat-value">42</div><div class="stat-title">Groups</div></div></div></div>
      <h2 class="section-title">Quote</h2><div class="module-demo"><blockquote class="wf-quote"><p>"Quote text."</p><cite>Attribution</cite></blockquote></div>
      <h2 class="section-title">Form steps and error summary</h2><div class="module-demo"><ol class="steps"><li class="done">Done</li><li class="current">Current</li><li>Next</li></ol><div class="error-summary"><h2>There is a problem</h2><ul><li><a href="#">Enter your email address</a></li></ul></div></div>
      <h2 class="section-title">Member banner and locked state</h2><div class="module-demo"><div class="member-banner"><div class="wf-placeholder avatar">Photo</div><div><p><strong>Name</strong> <span class="status-pill">Active</span></p><p class="meta-row"><span>Grade</span><span>Number</span></p></div><div><a href="#">Quick link</a></div></div><div class="locked" style="margin-top:1rem"><h3>For members only</h3><p>What this offers, and how to get access.</p></div></div>
      <h2 class="section-title">Embed and video placeholders</h2><div class="module-demo"><div class="grid grid-2"><div class="wf-placeholder ratio-16-9">Video 16:9, captions and pause control</div><div class="wf-placeholder embed">Interactive chart with table alternative</div></div></div>
'''
    out_body = body
    # library lives in modules/, uses the same page shell
    page('__library.html', 'Module library', out_body, '', [('Module library', None)])
    src = os.path.join(SRC, 'pages', '__library.html')
    os.makedirs(os.path.join(SRC, 'modules'), exist_ok=True)
    os.replace(src, os.path.join(SRC, 'modules', 'library.html'))
    PAGES_BUILT.remove('__library.html')


def requirements():
    """requirements/requirements.md: every Must and Should, and where it is shown."""
    import openpyxl
    xlsx = os.environ.get('RSS_REQS', '/mnt/user-data/outputs/RSS Website Requirements - ClerksWell review.xlsx')
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    rows = [r for r in list(wb.worksheets[0].iter_rows(values_only=True))[1:] if r and r[0]]
    # Where each requirement is referenced in the source pages
    where = {}
    for folder in ('pages', 'includes', 'modules'):
        base = os.path.join(SRC, folder)
        for f in sorted(os.listdir(base)):
            txt = open(os.path.join(base, f), encoding='utf-8').read()
            for a, b in re.findall(r'REQ-(\d{3})(?: to REQ-(\d{3}))?', txt):
                rng = range(int(a), int(b) + 1) if b else [int(a)]
                for i in rng:
                    where.setdefault(f'REQ-{i:03d}', set()).add(f)
    hubtxt = open(os.path.join(SRC, 'index.html'), encoding='utf-8').read()
    lines = ['# RSS website requirements: prototype trace', '',
             'Generated from the draft requirements catalogue (ClerksWell review copy). The catalogue is being revised; this trace will be regenerated from the revised set.', '',
             'Only Must and Should requirements are prototyped. "Shown in" lists the source files whose annotations reference the requirement; "not shown" usually means a platform, CMS or non-functional requirement with no front-end expression.', '']
    counts = {'shown': 0, 'not': 0}
    for prio in ('Must', 'Should'):
        lines += [f'## {prio}', '', '| Ref | Type | Area | Requirement | Shown in |', '| --- | --- | --- | --- | --- |']
        for r in rows:
            if str(r[6]).strip() != prio:
                continue
            ref = r[0]
            files = ', '.join(sorted(where.get(ref, [])))
            counts['shown' if files else 'not'] += 1
            desc = (r[2] or '').replace('|', '/').replace('\n', ' ')
            lines.append(f"| {ref} | {(r[1] or '')[:1]} | {(r[4] or '').replace('|', '/')} | {desc} | {files or 'not shown'} |")
        lines.append('')
    lines += ['## Could requirements referenced as "Proposed"', '']
    for r in rows:
        if str(r[6]).strip() == 'Could' and r[0] in where:
            lines.append(f"- {r[0]}: {(r[2] or '').replace(chr(10), ' ')} ({', '.join(sorted(where[r[0]]))})")
    os.makedirs(os.path.join(ROOT, 'requirements'), exist_ok=True)
    with open(os.path.join(ROOT, 'requirements', 'requirements.md'), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(lines) + '\n')
    return counts


if __name__ == '__main__':
    n_idx = write_data()
    pages_a.build()
    pages_b.build()
    pages_c.build()
    pages_d.build()
    hub()
    library()
    c = requirements()
    print(f'{len(PAGES_BUILT)} pages, {n_idx} search items; Must/Should shown: {c["shown"]}, not shown: {c["not"]}')
